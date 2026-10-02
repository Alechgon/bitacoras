"""
reportes.py — arma el Excel de reporte desde la agenda viva:
Resumen (metas), Gantt por día/bloque/técnico, detalle, hallazgos del grupo y
registro de cambios (aplazamientos, adelantos).

Uso: python reportes.py   -> reportes/Reporte_SOSER_AAAA-MM-DD.xlsx
"""
import os
from datetime import timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from nucleo import BASE, a_fecha, agenda, bloques, bonita, datos, db, habiles_entre, hoy, metas, nombre_tec, sumar_habiles

AZUL = PatternFill("solid", fgColor="1F4E78")
BLANCO = Font(bold=True, color="FFFFFF")
FINO = Border(*[Side(style="thin", color="BFBFBF")] * 4)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ARRIBA = Alignment(vertical="top", wrap_text=True)


def color(pts, clase=""):
    if "CORRECTIVO" in clase and pts >= 15:
        return "F8696B"
    if pts >= 13:
        return "FFB366"
    if pts >= 9:
        return "FFE699"
    return "C6EFCE"


def head(ws, fila, vals, anchos=None):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(fila, i, v)
        c.font, c.fill, c.border, c.alignment = BLANCO, AZUL, FINO, CENTRO
    if anchos:
        for i, w in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def celda(ws, f, c, v, fmt=None, fill=None, al=None):
    x = ws.cell(f, c, v)
    x.border = FINO
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    if al:
        x.alignment = al
    return x


def generar(dias=10):
    D = datos()
    E = D["E"]
    tecs = list(D["META"]["tecnicos"])
    con = db()
    h = hoy()
    fin = sumar_habiles(h, dias)
    wb = Workbook()

    # ------------------------------------------------ Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"SOSER San Pablo · Mantención PAE · {h.strftime('%d-%m-%Y')}"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A2"] = f"datos.js {D['META']['version']}"
    ws["A2"].font = Font(italic=True, color="808080")
    m = metas(con)
    head(ws, 4, ["Meta", "Cumplidos", "Total", "Avance", "Fecha meta", "Días hábiles restantes"])
    for i, (k, nombre) in enumerate([("junaeb", "REGLA 1 · JUNAEB en Datácora"),
                                     ("jardines", "REGLA 2 · Jardines con preventiva")], start=5):
        ok, tot, f = m[k]
        rest = sum(1 for _ in habiles_entre(h, a_fecha(f)))
        for j, v in enumerate([nombre, ok, tot, ok / tot if tot else 0, a_fecha(f), rest], start=1):
            celda(ws, i, j, v, "0%" if j == 4 else ("DD-MM-YYYY" if j == 5 else None))
    fila = 9
    head(ws, fila, ["Técnico", f"Visitas próximos {dias} días háb.", "Correctivos", "Visitas/Prev", "Bloques extra"])
    for t in tecs:
        fila += 1
        ag = agenda(con, h, fin, t)
        vals = [D["META"]["tecnicos"][t]["nombre"], len(ag), sum("CORRECTIVO" in x["clase"] for x in ag),
                sum("CORRECTIVO" not in x["clase"] for x in ag), sum(x["bloque"] > 3 for x in ag)]
        for j, v in enumerate(vals, start=1):
            celda(ws, fila, j, v)
    fila += 2
    sem = (h - timedelta(days=7)).isoformat()
    datos_h = [("Hallazgos del grupo últimos 7 días", con.execute(
                    "SELECT COUNT(*) FROM hallazgos WHERE recibido>=?", (sem,)).fetchone()[0]),
               ("Por confirmar (!es RBD)", con.execute(
                    "SELECT COUNT(*) FROM hallazgos WHERE estado='por_confirmar'").fetchone()[0]),
               ("Visitas aplazadas (programadas)", con.execute(
                    "SELECT COUNT(*) FROM tarjetas WHERE estado='programada' AND aplaz>0").fetchone()[0]),
               ("Visitas atrasadas sin marcar", con.execute(
                    "SELECT COUNT(*) FROM tarjetas WHERE estado='programada' AND fecha<?", (h.isoformat(),)
                ).fetchone()[0])]
    for k, v in datos_h:
        celda(ws, fila, 1, k)
        celda(ws, fila, 2, v)
        fila += 1
    for col, w in zip("ABCDEF", [38, 22, 14, 14, 14, 20]):
        ws.column_dimensions[col].width = w

    # ------------------------------------------------ Gantt día x bloque
    wg = wb.create_sheet("Agenda")
    head(wg, 1, ["Día", "Bloque"] + [D["META"]["tecnicos"][t]["nombre"].split()[0] for t in tecs],
         [12, 14] + [46] * len(tecs))
    r = 1
    for d in habiles_entre(h, fin):
        bl = bloques(d)
        extra = {x["bloque"] for x in agenda(con, d, d) if x["bloque"] not in bl}
        filas_dia = list(bl.items()) + [(b, "extra") for b in sorted(extra)]
        primera = r + 1
        for b, horario in filas_dia:
            r += 1
            celda(wg, r, 2, f"B{b} {horario}" if b in bl else "EXTRA", al=CENTRO)
            for j, t in enumerate(tecs, start=3):
                xs = [x for x in agenda(con, d, d, t) if x["bloque"] == b]
                txt = "\n".join(f"{E.get(x['rbd'], {}).get('nombre', x['rbd'])} ({x['pts']})\n{x['clase']}" +
                                (f" · {x['crit']}" if x["crit"] and "CORRECTIVO" in x["clase"] else "") for x in xs)
                celda(wg, r, j, txt, fill=color(max(x["pts"] for x in xs), xs[0]["clase"]) if xs else None, al=ARRIBA)
            wg.row_dimensions[r].height = 32
        wg.merge_cells(start_row=primera, end_row=r, start_column=1, end_column=1)
        c = wg.cell(primera, 1, bonita(d).upper())
        c.font, c.alignment, c.border = Font(bold=True), CENTRO, FINO
    wg.freeze_panes = "C2"

    # ------------------------------------------------ Detalle
    wd = wb.create_sheet("Detalle")
    head(wd, 1, ["Fecha", "Bloque", "Técnico", "RBD", "Establecimiento", "Comuna", "Inst.", "Supervisora", "Clase",
                 "Falla", "Pts", "Aplaz.", "Límite", "Origen", "Detalle"],
         [11, 7, 10, 9, 30, 15, 8, 18, 20, 8, 6, 7, 11, 8, 70])
    for i, x in enumerate(agenda(con, h, a_fecha(D["META"]["fin"])), start=2):
        e = E.get(x["rbd"], {})
        vals = [a_fecha(x["fecha"]), x["bloque"], nombre_tec(x["tec"]), x["rbd"], e.get("nombre"), e.get("comuna"),
                e.get("inst"), e.get("sup"), x["clase"], x["crit"], x["pts"], x["aplaz"] or "", a_fecha(x["limite"]),
                x["origen"], x["detalle"]]
        for j, v in enumerate(vals, start=1):
            celda(wd, i, j, v, "DD-MM-YYYY" if j in (1, 13) else None,
                  fill=color(x["pts"], x["clase"]) if j == 11 else None)
    wd.auto_filter.ref = wd.dimensions
    wd.freeze_panes = "A2"

    # ------------------------------------------------ Hallazgos
    wh = wb.create_sheet("Hallazgos WhatsApp")
    head(wh, 1, ["Recibido", "Supervisora", "RBD", "Establecimiento", "Problema", "Falla", "Pts", "Agendado", "Técnico",
                 "Bloque", "Estado", "Motor", "Mensaje"], [16, 20, 9, 30, 38, 8, 6, 11, 10, 7, 13, 18, 70])
    q = con.execute("""SELECT x.*, t.fecha AS f, t.tec, t.bloque FROM hallazgos x
                       LEFT JOIN tarjetas t ON t.id=x.tarjeta_id ORDER BY x.id DESC LIMIT 500""")
    for i, x in enumerate(q, start=2):
        vals = [x["recibido"], x["autor"], x["rbd"], E.get(x["rbd"], {}).get("nombre", "¿?"), x["problema"], x["crit"],
                x["pts"], a_fecha(x["f"]), nombre_tec(x["tec"]) if x["tec"] else "", x["bloque"], x["estado"],
                x["motor"], x["texto"]]
        for j, v in enumerate(vals, start=1):
            celda(wh, i, j, v, "DD-MM-YYYY" if j == 8 else None, al=ARRIBA if j == 13 else None,
                  fill=color(x["pts"] or 0, "CORRECTIVO") if j == 7 and x["pts"] else None)
    wh.auto_filter.ref = wh.dimensions
    wh.freeze_panes = "A2"

    # ------------------------------------------------ Mensajes (qué entendió el bot)
    wm = wb.create_sheet("Mensajes")
    head(wm, 1, ["Recibido", "Origen", "Autor", "Intención", "RBD", "Establecimiento", "Resultado", "Motor", "Texto"],
         [16, 8, 18, 14, 9, 28, 14, 18, 70])
    for i, x in enumerate(con.execute("SELECT * FROM mensajes ORDER BY id DESC LIMIT 800"), start=2):
        vals = [x["recibido"], x["origen"], x["autor"], x["intencion"], x["rbd"], E.get(x["rbd"], {}).get("nombre", ""),
                x["resultado"], x["motor"], x["texto"]]
        for j, v in enumerate(vals, start=1):
            celda(wm, i, j, v, al=ARRIBA if j == 9 else None)
    wm.auto_filter.ref = wm.dimensions
    wm.freeze_panes = "A2"

    # ------------------------------------------------ Observaciones
    wo = wb.create_sheet("Observaciones")
    head(wo, 1, ["Recibido", "Autor", "RBD", "Establecimiento", "Resumen", "Texto"], [16, 18, 9, 28, 40, 70])
    for i, x in enumerate(con.execute("SELECT * FROM observaciones ORDER BY id DESC LIMIT 500"), start=2):
        vals = [x["recibido"], x["autor"], x["rbd"], E.get(x["rbd"], {}).get("nombre", ""), x["resumen"], x["texto"]]
        for j, v in enumerate(vals, start=1):
            celda(wo, i, j, v, al=ARRIBA if j == 6 else None)

    # ------------------------------------------------ Cambios
    wc = wb.create_sheet("Cambios")
    head(wc, 1, ["Cuándo", "Evento"], [18, 120])
    for i, x in enumerate(con.execute("SELECT * FROM historial ORDER BY id DESC LIMIT 400"), start=2):
        celda(wc, i, 1, x["cuando"])
        celda(wc, i, 2, x["evento"])

    os.makedirs(os.path.join(BASE, "reportes"), exist_ok=True)
    ruta = os.path.join(BASE, "reportes", f"Reporte_SOSER_{h.isoformat()}.xlsx")
    wb.save(ruta)
    return ruta


if __name__ == "__main__":
    print("✅", generar())
