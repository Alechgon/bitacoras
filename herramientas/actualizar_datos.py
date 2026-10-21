#!/usr/bin/env python3
"""
actualizar_datos.py — pone al día datos.js (el panel y el bot leen este archivo).

Entradas:
  --datacora  export "Completadas" de Datácora (.xlsx)
  --registro  registro de bitácoras con el formato de Registro_Mantención (Libro1.xlsx):
              RBD, establecimiento, comuna, 1320, tipo estab, supervisora, Suc. RM, fecha,
              mes, técnico, tipo de trabajo, N° bitácora, descripción, ...
  --novedades JSON con los requerimientos de las supervisoras ya revisados y las visitas
              que todavía no están en Datácora (ver novedades_2026-10-20.json)

Qué hace, en orden:
  1. Reemplaza DC con el export (sin las cuentas de prueba ni CASA MATRIZ).
  2. Suma al historial (HIST) cada bitácora que no estaba, venga de Datácora, del registro o de
     un aviso.
  3. Agrupa las bitácoras en visitas (RBD + fecha + técnico) y las cruza con el plan:
     - la tarjeta planificada que esa visita cumple pasa a realizada, en la fecha y con el
       técnico reales;
     - si la visita no estaba en el plan, entra como tarjeta planificada y realizada
       ("como si hubiese estado planificada"); antes del 30-09 entra como histórica.
  4. Rehace CORR (requerimientos abiertos) con lo que quedó pendiente, y MENC con lo que
     dijeron las supervisoras.
  5. Replanifica desde --desde todo lo pendiente: JUNAEB sin bitácora en Datácora antes de la
     meta, correctivos de supervisoras en el bloque 3 dentro de 10 días hábiles, preventivas
     de jardines que tienen correctivos abiertos se adelantan para cerrar todo en una visita.
     Las tarjetas conservan su id, así el panel mantiene lo que marcaste a mano.
  6. Recalcula la ficha de cada establecimiento (cobertura, última visita, puntaje y su
     desglose) con los pesos del README.

Ojo: corre siempre sobre el datos.js ANTERIOR a la actualización (git checkout datos.js), no dos
veces seguidas sobre el mismo: la segunda vez contaría dos veces los aplazamientos.

Uso:
  python3 herramientas/actualizar_datos.py --datacora X.xlsx --registro Libro1.xlsx \
      --novedades herramientas/novedades_2026-10-20.json --desde 2026-10-21
"""
import argparse
import datetime as dt
import json
import math
import os
import re
import unicodedata
from collections import defaultdict

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(AQUI)

AUTORES = {
    "Annahi": "Annahi Uribe( Super Soser)",
    "Stephania": "Stephania Fernandez Supp Soser",
    "Carla": "Carla Espinoza (SUP Santiago Soser)",
    "Juan Carlos": "Juan Carlos",
}
SEV = {"GAS": 6, "FRIO": 4, "AGUA": 4, "ELEC": 3, "EQUIPO": 2, "OTRO": 1}
TIPO_DC = {"Emergencia": "Correctiva", "Plan Preventivo Mantención": "Preventiva",
           "Mutualidad": "Mutualidad", "Acta": "Acta"}
TEC_NOMBRE = {"RODRIGO": "Rodrigo Martinez", "CAMILO": "Camilo Santis"}


# ============================================================ utilidades
def iso(x):
    if hasattr(x, "strftime"):
        return x.strftime("%Y-%m-%d")
    return str(x) if x else ""


def F(s):
    return dt.date.fromisoformat(s)


def fc(s):
    """28-10"""
    return F(s).strftime("%d-%m")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def tec_de(s):
    s = norm(s)
    if "rodrigo" in s:
        return "RODRIGO"
    if "camilo" in s:
        return "CAMILO"
    return None


def inst_de(s):
    s = norm(s)
    return "Junji" if "junji" in s else "Integra" if "integra" in s else "Junaeb"


def es_prev(t):
    return bool(re.search(r"plan preventivo|preventiv", t or "", re.I))


def cargar_js(ruta):
    t = open(ruta, encoding="utf-8").read()
    i, j = t.index("{"), t.rindex("}")
    return json.loads(t[i:j + 1])


def guardar_js(ruta, D, cuando):
    cab = (f"// SOSER San Pablo — datos actualizados el {cuando} con herramientas/actualizar_datos.py. "
           "No editar a mano.\n")
    with open(ruta, "w", encoding="utf-8") as f:
        f.write(cab + "window.SOSER = " + json.dumps(D, ensure_ascii=False, indent=1) + ";\n")


# ============================================================ lectura de fuentes
def leer_datacora(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True)["Completadas"]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c or "") for c in filas[0]]
    col = {n: cab.index(n) for n in cab}
    out = []
    for r in filas[1:]:
        g = lambda n: r[col[n]] if n in col else None
        folio = g("Folio bitácora")
        tec = tec_de(g("Técnico"))
        rbd = str(g("RBD") or "")
        if not folio or not tec or not rbd.isdigit() or int(rbd) == 1:
            continue  # cuentas de prueba (Prueba stgo, Google play) y CASA MATRIZ
        tipo = TIPO_DC.get(g("Tipo de visita"), g("Tipo de visita") or "Correctiva")
        inst = inst_de(g("Institución"))
        out.append({
            "fecha": iso(g("Fecha de cierre")), "fplan": iso(g("Fecha planificada")),
            "rbd": int(rbd), "tipo": tipo, "folio": str(folio), "tec": TEC_NOMBRE[tec], "_tec": tec,
            "estab": g("Establecimiento") or "", "inst": inst,
            "det": (g("Descripción / observación") or "").strip(),
            # para JUNAEB vale cualquier visita; en jardines solo la preventiva
            "cuenta": inst == "Junaeb" or tipo == "Preventiva",
            "revision": g("Revisión del plan") or "",
            "_tarea": g("ID tarea") or "",
        })
    out.sort(key=lambda d: (d["fecha"], int(d["folio"])), reverse=True)
    return out


def leer_registro(ruta):
    ws = openpyxl.load_workbook(ruta, data_only=True).active
    out = []
    for r in ws.iter_rows(values_only=True):
        # la tabla puede venir corrida: se busca la primera celda que sea un RBD
        k = next((i for i, c in enumerate(r) if isinstance(c, int) and c > 1000), None)
        if k is None or len(r) < k + 13 or not hasattr(r[k + 7], "strftime"):
            continue
        fila = r[k:k + 16]
        tt = norm(fila[10])
        desc = (fila[12] or "").strip()
        tipo = "Preventiva" if ("plan" in tt or "prevent" in tt) else "Correctiva"
        parte = "SC" if re.search(r"(\(\s*sc\s*\)|\bsc\s*$)", desc, re.I) else "JI"
        out.append({"rbd": fila[0], "estab": str(fila[1]).replace("\xa0", " ").strip(),
                    "fecha": iso(fila[7]), "tec": tec_de(fila[9]), "tipo": tipo,
                    "folio": str(fila[11] or "s/n"), "det": desc, "parte": parte})
    return out


# ============================================================ geografía (para elegir técnico)
def coords(e, E, base_comuna):
    if e.get("lat") and e.get("lon"):
        return e["lat"], e["lon"]
    return base_comuna.get(e["comuna"], (-33.46, -70.66))


def km(a, b):
    return math.hypot((a[0] - b[0]) * 111.0, (a[1] - b[1]) * 93.0)


# ============================================================ principal
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--datos", default=os.path.join(RAIZ, "datos.js"))
    ap.add_argument("--datacora", required=True)
    ap.add_argument("--registro")
    ap.add_argument("--novedades", required=True)
    ap.add_argument("--desde", required=True, help="primer día a replanificar (AAAA-MM-DD)")
    ap.add_argument("--version")
    ap.add_argument("--firmas-de", nargs="*", default=[],
                    help="datos.js de versiones anteriores que alguien pudo dejar guardadas en su navegador")
    ap.add_argument("--reporte", default=None)
    a = ap.parse_args()

    D = cargar_js(a.datos)
    NOV = json.load(open(a.novedades, encoding="utf-8"))
    M = D["META"]
    # Firmas de cada tarjeta tal como venía en las versiones anteriores. El panel las usa para
    # distinguir lo que tú cambiaste a mano de lo que solo venía del plan viejo (ver fusionar()).
    previas = {M["version"]: firmas(D)}
    for ruta in a.firmas_de:
        X = cargar_js(ruta)
        previas.setdefault(X["META"]["version"], firmas(X))
    version_entrada = M["version"]
    E = {e["rbd"]: e for e in D["ESTAB"]}
    DESDE = a.desde
    CORTE = NOV["corte"]
    INICIO = M["inicio"]
    rep = []                                     # líneas del reporte
    avisos = []

    # ---------------------------------------------------------------- 1. Datácora
    dc = leer_datacora(a.datacora)
    for d in dc:                                  # correcciones tuyas al tipo que quedó en Datácora
        r = NOV.get("reclasificar_datacora", {}).get(d["folio"])
        if r:
            d["tipo"] = r["tipo"]
            d["cuenta"] = d["inst"] == "Junaeb" or d["tipo"] == "Preventiva"
            d["det"] = (d["det"] + "  [" + r["nota"] + "]").strip()
            rep.append(f"  Folio {d['folio']} {d['estab']}: queda como {d['tipo']} ({r['nota']})")
    separadas = {f: i for i, g in enumerate(NOV.get("visitas_separadas", [])) for f in g}
    fuera = sorted({d["rbd"] for d in dc if d["rbd"] not in E})
    if fuera:
        avisos.append(f"Datácora trae RBD que no están en la sucursal: {fuera}")
    dc = [d for d in dc if d["rbd"] in E]
    viejos = {d["folio"] for d in D["DC"]}
    nuevos_dc = [d for d in dc if d["folio"] not in viejos]
    perdidos = viejos - {d["folio"] for d in dc}
    if perdidos:
        avisos.append(f"Folios que estaban en el export anterior y ya no aparecen: {sorted(perdidos)}")
    dup = defaultdict(list)
    for d in dc:
        if d["_tarea"]:
            dup[d["_tarea"]].append(d["folio"])
    dobles = [v for v in dup.values() if len(v) > 1 and not all(f in separadas for f in v)]
    D["DC"] = [{k: v for k, v in d.items() if not k.startswith("_")} for d in dc]
    rep.append(f"Datácora: {len(dc)} bitácoras válidas ({len(nuevos_dc)} nuevas desde el export anterior).")
    for v in dobles:
        rep.append(f"  Ojo: la misma tarea se envió dos veces en Datácora (folios {' y '.join(v)}).")

    # ---------------------------------------------------------------- 2. bitácoras → HIST
    reg = leer_registro(a.registro) if a.registro else []
    bits = []                                    # todas las bitácoras, normalizadas
    for d in dc:
        bits.append({"fecha": d["fecha"], "rbd": d["rbd"], "tec": d["_tec"], "tipo": d["tipo"],
                     "parte": "JI", "folio": "D-" + d["folio"], "num": d["folio"],
                     "det": d["det"], "fuente": "Datácora"})
    for x in reg:
        if x["rbd"] not in E or not x["tec"]:
            avisos.append(f"Registro: fila sin RBD de la sucursal o sin técnico: {x}")
            continue
        bits.append({"fecha": x["fecha"], "rbd": x["rbd"], "tec": x["tec"], "tipo": x["tipo"],
                     "parte": x["parte"], "folio": x["folio"], "num": x["folio"],
                     "det": x["det"], "fuente": "Registro"})
    for v in NOV.get("visitas_sin_datacora", []):
        bits.append({"fecha": v["fecha"], "rbd": v["rbd"], "tec": v["tec"], "tipo": v["tipo"],
                     "parte": v.get("parte", "JI"), "folio": v.get("folio", "s/n"), "num": v.get("folio", "s/n"),
                     "det": v["det"], "fuente": "Aviso (sin Datácora)"})

    # cotejo de fechas: registro contra Datácora
    confirmados = set(NOV.get("registro_confirmado", []))
    for x in [b for b in bits if b["fuente"] == "Registro" and b["folio"] not in confirmados]:
        cand = [d for d in dc if d["rbd"] == x["rbd"] and abs((F(d["fecha"]) - F(x["fecha"])).days) <= 6]
        mismo = [d for d in cand if d["fecha"] == x["fecha"]]
        if mismo:
            continue
        if cand:
            c = min(cand, key=lambda d: abs((F(d["fecha"]) - F(x["fecha"])).days))
            rep.append(f"  Revisar fecha: registro {fc(x['fecha'])} folio {x['folio']} {E[x['rbd']]['nombre']} no tiene par en "
                       f"Datácora ese día; la más cercana es {fc(c['fecha'])} folio {c['folio']} («{c['det'][:40]}»). "
                       f"Si es el mismo trabajo, corrige la fecha en el registro; si no, está bien.")
        else:
            x["solo_papel"] = True

    hist_folios = {str(h["folio"]) for h in D["HIST"]}
    hist_vis = {(h["rbd"], h["fecha"], h["tec"]) for h in D["HIST"]}
    n_hist = 0
    for b in bits:
        if b["folio"] in hist_folios and b["folio"] not in ("s/n", ""):
            continue
        # completa registros que estaban sin folio (p.ej. preventiva "aún no subida")
        hueco = next((h for h in D["HIST"] if h["rbd"] == b["rbd"] and h["fecha"] == b["fecha"]
                      and not h.get("folio") and h["tipo"] == b["tipo"]), None)
        if hueco:
            hueco["folio"] = b["folio"]
            hueco["detalle"] = b["det"] or hueco["detalle"]
            hueco["tec"] = b["tec"]
            continue
        D["HIST"].append({"fecha": b["fecha"], "rbd": b["rbd"], "tipo": b["tipo"], "parte": b["parte"],
                          "folio": b["folio"], "tec": b["tec"],
                          "detalle": (b["det"] + ("  [Datácora]" if b["fuente"] == "Datácora" else "")).strip(),
                          "fuente": b["fuente"]})
        hist_folios.add(b["folio"])
        n_hist += 1
    D["HIST"].sort(key=lambda h: (h["fecha"], h["rbd"]))
    rep.append(f"Historial: {n_hist} bitácoras nuevas (Datácora + registro + avisos).")

    # ---------------------------------------------------------------- 3. visitas
    vis = {}
    for b in sorted(bits, key=lambda b: b["fecha"]):
        # dos folios marcados como visitas distintas no se juntan aunque sean del mismo día
        k = (b["rbd"], b["fecha"], b["tec"], b["num"] if b["num"] in separadas else "")
        v = vis.setdefault(k, {"rbd": b["rbd"], "fecha": b["fecha"], "tec": b["tec"], "tipos": set(),
                               "partes": set(), "folios": [], "det": [], "fuentes": set()})
        v["tipos"].add(b["tipo"])
        if b["tipo"] == "Preventiva":
            v["partes"].update(x.strip() for x in b["parte"].split("+") if x.strip())
        if b["num"] not in v["folios"]:
            v["folios"].append(b["num"])
        if b["det"] and b["det"] not in v["det"]:
            v["det"].append(b["det"])
        v["fuentes"].add(b["fuente"])
    visitas = sorted(vis.values(), key=lambda v: (v["fecha"], v["rbd"]))

    def tipo_vis(v):
        if "Preventiva" in v["tipos"]:
            return "Preventiva"
        if "Mutualidad" in v["tipos"]:
            return "Mutualidad"
        return "Correctiva"

    def resumen(v, n=150):
        t = " / ".join(x.replace("\n", " ").strip() for x in v["det"])
        t = re.sub(r"\s+", " ", t)
        return t[:n] + ("…" if len(t) > n else "")

    P = D["PLAN"]
    porid = {p["id"]: p for p in P}
    n_c = max(int(p["id"][1:]) for p in P if p["id"].startswith("C"))
    n_p = max(int(p["id"][1:]) for p in P if p["id"].startswith("P"))

    def nuevo_id(pref):
        nonlocal n_c, n_p
        if pref == "C":
            n_c += 1
            return f"C{n_c:04d}"
        n_p += 1
        return f"P{n_p:04d}"

    # partes de preventiva ya hechas en el 2° semestre (antes de cruzar)
    def partes_hechas(hasta=None):
        h = defaultdict(set)
        for p in P:
            if p["estado"] == "realizada" and p["fecha"] >= M["corteS2"] and es_prev(p.get("tipoReal")):
                if hasta and p["fecha"] > hasta:
                    continue
                for x in (p.get("parte") or "JI").split("+"):
                    h[p["rbd"]].add(x)
        return h

    # 3a. antes del inicio del plan: visitas que faltan en el histórico
    hist_cards = defaultdict(list)
    for p in P:
        if p["origen"] == "historico":
            hist_cards[(p["rbd"], p["fecha"], p["tec"])].append(p)
    n_hc = 0
    for v in visitas:
        if v["fecha"] >= INICIO:
            continue
        ya = hist_cards.get((v["rbd"], v["fecha"], v["tec"]), [])
        tv = tipo_vis(v)
        partes_ya = {p.get("parte") or "JI" for p in ya if es_prev(p.get("tipoReal"))}
        falta_partes = sorted(v["partes"] - partes_ya) if tv == "Preventiva" else []
        if ya and not falta_partes:
            for p in ya:                                  # completa el folio si faltaba
                if not p.get("folio"):
                    p["folio"] = ", ".join(v["folios"])
            continue
        for parte in (["+".join(falta_partes)] if falta_partes else ["JI"]):
            card = {"id": nuevo_id("C"), "fase": "HIST", "fecha": v["fecha"], "tec": v["tec"], "bloque": 1, "hora": "",
                    "rbd": v["rbd"], "clase": "PREVENTIVA" if tv == "Preventiva" else "CORRECTIVO URGENTE",
                    "detalle": f"{tv} ({parte}) · bitácora {', '.join(v['folios'])} · {resumen(v, 110)}",
                    "tipoReal": tv, "parte": parte, "fuente": " + ".join(sorted(v["fuentes"])),
                    "pts": E[v["rbd"]]["pts"], "crit": "", "folio": ", ".join(v["folios"]),
                    "estado": "realizada", "aplaz": 0, "origen": "historico"}
            P.append(card)
            hist_cards[(v["rbd"], v["fecha"], v["tec"])].append(card)
            n_hc += 1
    rep.append(f"Histórico (antes del {fc(INICIO)}): {n_hc} tarjetas nuevas.")

    # 3b. desde el inicio del plan: cruzar con las tarjetas planificadas
    corr_prev = NOV.get("corr_previos", {})
    REQ = NOV["requerimientos"]
    pend_tec = set()      # RBD con requerimientos que necesitan un técnico
    for r in REQ:
        if r["estado"] == "pendiente" or (r["estado"] == "parcial" and not r.get("cierre")):
            pend_tec.add(r["rbd"])
    for k, c in corr_prev.items():
        if c["estado"] in ("pendiente", "parcial"):
            pend_tec.add(int(k))

    pend_cards = defaultdict(list)
    for p in P:
        if p["origen"] == "plan" and p["estado"] == "programada":
            pend_cards[p["rbd"]].append(p)
    for l in pend_cards.values():
        l.sort(key=lambda p: p["fecha"])

    n_hechas, n_agreg = 0, 0
    prev_plan = partes_hechas()
    usadas = set()
    # visitas que ya tenían tarjeta realizada (el histórico del export anterior llega al 30-09)
    ya_real = {}
    for p in P:
        if p["estado"] == "realizada":
            ya_real.setdefault((p["rbd"], p["fecha"], p["tec"]), p)

    def cerrado_previo(rbd):
        return corr_prev.get(str(rbd), {}).get("estado", "cerrado") == "cerrado"

    def marcar_hecha(p, v, parte=None):
        nonlocal n_hechas
        antes = p["fecha"]
        tv = tipo_vis(v)
        p["estado"] = "realizada"
        p["fecha"], p["tec"] = v["fecha"], v["tec"]
        p["folio"] = ", ".join(v["folios"])
        p["tipoReal"] = tv
        if tv == "Preventiva":
            p["parte"] = parte or "+".join(sorted(v["partes"])) or "JI"
            if "CORRECTIVO" in p["clase"]:
                p["clase"] = "PREVENTIVA + CORRECTIVO"
            elif p["clase"] != "PREVENTIVA":
                p["clase"] = "PREVENTIVA"
        nota = f"Hecha el {fc(v['fecha'])} por {v['tec'].title()} · bitácora {p['folio']}"
        if antes != v["fecha"]:
            nota += f" (estaba para el {fc(antes)})"
        p["detalle"] = nota + " · " + resumen(v, 120) + "  ||  " + p["detalle"]
        n_hechas += 1

    for v in visitas:
        if v["fecha"] < INICIO:
            continue
        e = E[v["rbd"]]
        tv = tipo_vis(v)
        cards = [p for p in pend_cards[v["rbd"]] if p["id"] not in usadas]
        elegida = None
        if e["inst"] == "Junaeb":
            elegida = next((p for p in cards if p["clase"].startswith("VISITA")), None)
        else:
            if tv == "Preventiva":
                elegida = next((p for p in cards if p["clase"].startswith("PREVENTIVA")), None)
                if elegida:
                    nec = {"JI", "SC"} if e["nbit"] == 2 else {"JI"}
                    tiene = prev_plan[v["rbd"]] | v["partes"]
                    if not nec <= tiene:
                        elegida = None          # falta otra parte: la tarjeta sigue abierta
            if not elegida and v["rbd"] not in pend_tec and cerrado_previo(v["rbd"]):
                elegida = next((p for p in cards if p["clase"] == "CORRECTIVO URGENTE"), None)
        previa = ya_real.get((v["rbd"], v["fecha"], v["tec"]))
        # una visita = una tarjeta; si dejó JI y SC, la tarjeta lleva parte "JI+SC"
        if previa:
            # la visita ya estaba como realizada: se le suman los folios y la planificada se absorbe
            fol = [f for f in v["folios"] if f not in str(previa.get("folio", ""))]
            if fol:
                previa["folio"] = ", ".join(x for x in [str(previa.get("folio") or "").replace("D-", "")] + fol if x)
            if elegida:
                usadas.add(elegida["id"])
                elegida["estado"] = "anulada"
                elegida["absorbidaPor"] = previa["id"]
                elegida["detalle"] = (f"Hecha el {fc(v['fecha'])} por {v['tec'].title()} (bitácora {previa['folio']}), "
                                      f"ver {previa['id']}  ||  " + elegida["detalle"])
                n_hechas += 1
        elif elegida:
            usadas.add(elegida["id"])
            marcar_hecha(elegida, v)
        else:
            parte = "+".join(sorted(v["partes"])) if tv == "Preventiva" else ""
            clase = ("PREVENTIVA" if tv == "Preventiva" else "VISITA" if tv == "Mutualidad"
                     else "CORRECTIVO URGENTE")
            card = {"id": nuevo_id("P"), "fase": "F1", "fecha": v["fecha"], "tec": v["tec"], "bloque": 3,
                    "hora": "", "rbd": v["rbd"], "clase": clase,
                    "detalle": (f"{tv}{' (' + parte + ')' if parte else ''} · bitácora {', '.join(v['folios'])} · "
                                f"{resumen(v)}"),
                    "tipoReal": tv, "parte": parte or "JI", "pts": e["pts"], "crit": "",
                    "folio": ", ".join(v["folios"]), "estado": "realizada", "aplaz": 0, "origen": "plan",
                    "fuente": " + ".join(sorted(v["fuentes"]))}
            P.append(card)
            n_agreg += 1
        for x in v["partes"]:
            prev_plan[v["rbd"]].add(x)
    rep.append(f"Plan: {n_hechas} tarjetas planificadas quedaron realizadas; "
               f"{n_agreg} visitas que no estaban en el plan entraron como planificadas y realizadas.")

    # tarjetas de preventiva que quedaron cumplidas por bitácoras anteriores al plan
    prev_tot = partes_hechas()
    for v in visitas:
        for x in v["partes"]:
            prev_tot[v["rbd"]].add(x)
    for p in P:
        if p["origen"] == "plan" and p["estado"] == "programada" and p["clase"].startswith("PREVENTIVA"):
            e = E[p["rbd"]]
            nec = {"JI", "SC"} if e["nbit"] == 2 else {"JI"}
            falta = nec - prev_tot[p["rbd"]]
            if not falta:
                ult = max((c for c in P if c["rbd"] == p["rbd"] and c["estado"] == "realizada"
                           and es_prev(c.get("tipoReal")) and c["fecha"] >= M["corteS2"]),
                          key=lambda c: c["fecha"], default=None)
                p["estado"] = "anulada"
                if ult:
                    p["absorbidaPor"] = ult["id"]
                    p["detalle"] = (f"Ya no hace falta: la preventiva quedó completa el {fc(ult['fecha'])} "
                                    f"(bitácora {ult.get('folio', '')})  ||  " + p["detalle"])
                else:
                    p["detalle"] = "Ya no hace falta: la preventiva del 2° semestre está completa  ||  " + p["detalle"]
            else:
                p["parte"] = "+".join(sorted(falta))
                if p["parte"] == "SC":
                    p["detalle"] = re.sub(r"^PREVENTIVA obligatoria — [^·]*", "PREVENTIVA obligatoria — ya tiene la parte JI, falta levantar SC ", p["detalle"])

    # ---------------------------------------------------------------- 4. requerimientos → CORR y MENC
    hoy = NOV["hoy"]
    viejos_corr = {c["rbd"]: c for c in D["CORR"]}
    abiertos = defaultdict(list)            # rbd -> requerimientos abiertos
    cerrados = []
    for r in REQ:
        if r["estado"] in ("cerrado", "info"):
            cerrados.append(r)
        elif r["estado"] == "parcial" and r.get("cierre"):
            cerrados.append(r)
        else:
            abiertos[r["rbd"]].append(r)
    for rbd, c in viejos_corr.items():
        cp = corr_prev.get(str(rbd), {})
        st = cp.get("estado", "pendiente")
        if st == "cerrado":
            continue
        texto = cp.get("texto") or re.sub(r"\s+—\s+prioridad declarada.*$", "", c["detalle"])
        # si el mismo reclamo ya viene en la revisión del chat, no se repite
        if any(norm(texto)[:25] in norm(r["texto"]) for r in REQ if r["rbd"] == rbd):
            continue
        abiertos[rbd].insert(0, {"id": "C-" + str(rbd), "fecha": c["freporte"], "rbd": rbd, "crit": c["crit"],
                                 "prio": c.get("prio", False), "texto": texto,
                                 "estado": st, "nota": cp.get("nota", ""), "_viejo": True})

    def quien(rs):
        t = {r["estado"] for r in rs}
        if t & {"pendiente", "parcial"}:
            return "tecnico"
        return "proveedor" if "proveedor" in t else "admin"

    CORR = []
    for rbd, rs in abiertos.items():
        rs.sort(key=lambda r: r["fecha"])
        crit = max((r["crit"] for r in rs if r["estado"] in ("pendiente", "parcial")), key=lambda c: SEV[c],
                   default=max((r["crit"] for r in rs), key=lambda c: SEV[c]))
        q = quien(rs)
        estados = {r["estado"] for r in rs}
        estado = ("PENDIENTE" if "pendiente" in estados else "PARCIAL" if "parcial" in estados
                  else "PROVEEDOR" if "proveedor" in estados else "ADMIN")
        txt = []
        for r in rs:
            t = r["texto"]
            if r["estado"] in ("proveedor", "admin"):
                t += " [" + ("proveedor" if r["estado"] == "proveedor" else "gestión") + "]"
            if r.get("nota") and not r.get("_viejo"):
                t += f" ({r['nota']})"
            elif r.get("nota"):
                t += f" — {r['nota']}"
            txt.append((fc(r["fecha"]) + " " + t) if not r.get("_viejo") else t)
        prio = any(r.get("prio") for r in rs)
        fre = min(r["fecha"] for r in rs if r["estado"] in ("pendiente", "parcial")) if q == "tecnico" else rs[0]["fecha"]
        CORR.append({"rbd": rbd, "estado": estado, "detalle": " · ".join(txt), "crit": crit, "freporte": fre,
                     "dias": (F(hoy) - F(fre)).days, "prio": prio, "pts": 0, "agendado": None, "quien": q,
                     "reqs": [r["id"] for r in rs]})
    D["CORR_CERRADOS"] = [{"id": r["id"], "fecha": r["fecha"], "rbd": r["rbd"], "texto": r["texto"],
                           "estado": r["estado"], "cierre": r.get("cierre"), "folio": r.get("folio"),
                           "nota": r.get("nota", "")} for r in cerrados]

    # MENC: lo que dijeron las supervisoras (sin repetir lo que ya estaba)
    ya_m = {(m["fecha"], m["rbd"], norm(m["texto"])[:40]) for m in D["MENC"]}
    ya_m2 = {(m["fecha"], m["rbd"]) for m in D["MENC"]}
    n_m = 0
    for r in REQ:
        k = (r["fecha"], r["rbd"], norm(r["texto"])[:40])
        if k in ya_m or ((r["fecha"], r["rbd"]) in ya_m2 and r["fecha"] < "2026-09-30"):
            continue
        D["MENC"].append({"fecha": r["fecha"], "rbd": r["rbd"], "autor": AUTORES.get(r["autor"], r["autor"]),
                          "texto": r["texto"] + (f"  →  {r['estado'].upper()}" +
                                                 (f" {fc(r['cierre'])} (folio {r['folio']})" if r.get("cierre") else "")
                                                 if r["estado"] != "info" else "")})
        n_m += 1
    D["MENC"].sort(key=lambda m: (m["fecha"], m["rbd"]))
    rep.append(f"Supervisoras: {len(REQ)} requerimientos revisados · {len(cerrados)} cerrados o solo aviso · "
               f"{sum(1 for c in CORR if c['quien'] == 'tecnico')} establecimientos con trabajo pendiente para técnico · "
               f"{sum(1 for c in CORR if c['quien'] != 'tecnico')} solo de proveedor o gestión. {n_m} menciones nuevas.")

    # ---------------------------------------------------------------- 6 (antes de planificar). fichas
    recalcular_fichas(D, E, CORR, M)
    corr_de = {c["rbd"]: c for c in CORR}
    for c in CORR:
        c["pts"] = E[c["rbd"]]["pts"]

    # ---------------------------------------------------------------- 5. replanificar
    dias = [d for d in M["habiles"] if d >= DESDE]
    vie = lambda d: F(d).weekday() == 4
    ocup = defaultdict(set)              # (fecha, tec) -> bloques ocupados
    sitios = defaultdict(list)           # (fecha, tec) -> coords asignadas
    base_comuna = {}
    for com in {e["comuna"] for e in D["ESTAB"]}:
        pts = [(e["lat"], e["lon"]) for e in D["ESTAB"] if e["comuna"] == com and e.get("lat")]
        base_comuna[com] = (sum(p[0] for p in pts) / len(pts), sum(p[1] for p in pts) / len(pts))
    BASE = {"RODRIGO": base_comuna.get("Estación Central", (-33.46, -70.69)),
            "CAMILO": base_comuna.get("Santiago", (-33.46, -70.65))}

    # cumplimiento de la meta JUNAEB con lo que hay hoy
    okJ = {d["rbd"] for d in D["DC"] if d["cuenta"]}
    okJ |= {p["rbd"] for p in P if p["estado"] == "realizada" and p["fecha"] >= CORTE}
    faltaJ = {e["rbd"] for e in D["ESTAB"] if e["inst"] == "Junaeb" and e["rbd"] not in okJ}

    # tarjetas que se quedan donde están: preventivas de noviembre y lo ya realizado
    tec_corr = {c["rbd"] for c in CORR if c["quien"] == "tecnico"}
    fijas = [p for p in P if p["estado"] == "programada" and p["origen"] == "plan" and p["fecha"] >= DESDE
             and p["clase"].startswith("PREVENTIVA") and p["rbd"] not in tec_corr]
    for p in fijas:
        ocup[(p["fecha"], p["tec"])].add(p["bloque"])
        sitios[(p["fecha"], p["tec"])].append(coords(E[p["rbd"]], E, base_comuna))

    trabajos = []
    pendientes = [p for p in P if p["estado"] == "programada" and p["origen"] == "plan" and p not in fijas]
    por_rbd = defaultdict(list)
    for p in pendientes:
        por_rbd[p["rbd"]].append(p)
    rbds = set(por_rbd) | {c["rbd"] for c in CORR if c["quien"] == "tecnico"} | faltaJ
    for rbd in rbds:
        e = E[rbd]
        cs = sorted(por_rbd.get(rbd, []), key=lambda p: (not p["clase"].startswith("PREVENTIVA"), p["fecha"]))
        c = corr_de.get(rbd)
        tec_req = c and c["quien"] == "tecnico"
        prev_card = next((p for p in cs if p["clase"].startswith("PREVENTIVA")), None)
        otra = next((p for p in cs if not p["clase"].startswith("PREVENTIVA")), None)
        junaeb_meta = rbd in faltaJ
        if not (tec_req or junaeb_meta or cs):
            continue
        if not tec_req and not junaeb_meta:
            # visita de octubre que ya no hace falta para la meta: seguimiento después del plazo
            for p in cs:
                if p["clase"].startswith("PREVENTIVA"):
                    continue
                if e["inst"] != "Junaeb":
                    p["estado"] = "anulada"
                    hecha = max((x for x in P if x["rbd"] == rbd and x["estado"] == "realizada"
                                 and x["fecha"] >= INICIO), key=lambda x: x["fecha"], default=None)
                    if c and c["quien"] != "tecnico":
                        p["detalle"] = ("Sale del plan de técnicos: " +
                                        ("lo ve el proveedor" if c["quien"] == "proveedor" else "es gestión tuya") +
                                        ", sigue en Pendientes  ||  " + p["detalle"])
                    elif hecha:
                        p["absorbidaPor"] = hecha["id"]
                        p["detalle"] = (f"Resuelto en la visita del {fc(hecha['fecha'])} ({hecha['id']}, bitácora "
                                        f"{hecha.get('folio', '')})  ||  " + p["detalle"])
                    else:
                        p["detalle"] = "Sale del plan: supervisión lo dio por resuelto  ||  " + p["detalle"]
                    continue
                trabajos.append({"rbd": rbd, "card": p, "absorbe": [], "tipo": "seguimiento", "need": "plan",
                                 "deadline": M["fin"], "prio": False, "crit": "", "pts": e["pts"],
                                 "fre": M["metaJunaeb"]})
            continue
        principal = prev_card if (prev_card and tec_req) else (otra or prev_card)
        absorbe = [p for p in cs if p is not principal and not (p is prev_card and not tec_req)]
        if junaeb_meta and tec_req:
            tipo, need = "visita+corr", "plan"
        elif junaeb_meta:
            tipo, need = "visita", "plan"
        elif principal is prev_card and prev_card is not None:
            tipo, need = "prev+corr", "plan"
        else:
            tipo, need = "corr", "emerg"
        fre = c["freporte"] if tec_req else M["metaJunaeb"]
        dl = sumar_habiles(fre, 10, M) if tec_req else M["metaJunaeb"]
        if tec_req:                            # lo ya vencido: dentro de los próximos 3 días hábiles
            dl = max(dl, sumar_habiles(DESDE, 2, M))
        if junaeb_meta:
            dl = min(dl, M["metaJunaeb"])
        if rbd == 947118:                      # visita de JUNJI el 28-10
            dl = min(dl, "2026-10-27")
        trabajos.append({"rbd": rbd, "card": principal, "absorbe": absorbe, "tipo": tipo, "need": need,
                         "deadline": dl, "prio": bool(c and c["prio"]), "crit": c["crit"] if c else "",
                         "pts": e["pts"], "fre": fre})

    # REGLA N°1 manda: primero los JUNAEB que faltan en Datácora (bloque 1 de cada técnico),
    # después los correctivos (bloque 3, y el 2 si no alcanza) y al final los seguimientos.
    for t in trabajos:
        if t["tipo"] in ("visita", "visita+corr"):
            t["need"] = "b1"
    orden = {"visita": 0, "visita+corr": 0, "corr": 1, "prev+corr": 1, "seguimiento": 3}
    trabajos.sort(key=lambda t: (orden[t["tipo"]],
                                 (-t["pts"],) if orden[t["tipo"]] == 0 else
                                 (max(t["deadline"], DESDE), t["crit"] != "GAS", not t["prio"],
                                  -SEV.get(t["crit"], 0), t["fre"], -t["pts"])))

    def libres(d, tec, need):
        if need == "b1":
            bl = [1]
        elif need == "plan":
            bl = [1] if vie(d) else [1, 2]
        else:
            bl = [2] if vie(d) else [3]
        return [b for b in bl if b not in ocup[(d, tec)]]

    def elegir_tec(t, d, need):
        e = E[t["rbd"]]
        cand = [x for x in ("RODRIGO", "CAMILO") if libres(d, x, need)]
        if not cand:
            return None
        c = corr_de.get(t["rbd"])
        gasfit = c and c["quien"] == "tecnico" and re.search(r"calef|horno|anafe|fog[oó]n|cocinilla|gas", norm(c["detalle"]))
        pref = ("CAMILO" if t["crit"] == "GAS" or (t["crit"] == "AGUA" and gasfit)
                else "RODRIGO" if t["crit"] == "ELEC" else None)
        if pref in cand:
            return pref
        aqui = coords(e, E, base_comuna)

        def costo(x):
            s = sitios[(d, x)]
            ref = min((km(aqui, y) for y in s), default=km(aqui, BASE[x]) + 1.5)
            return (ref, len(s))
        return min(cand, key=costo)

    reprog = []
    for t in trabajos:
        lugar = None
        if t["need"] == "b1":
            intentos = [("b1", True), ("plan", True), ("emerg", True), ("plan", False)]
        elif t["tipo"] == "seguimiento":
            intentos = [("plan", True), ("plan", False)]
        else:
            otra = "emerg" if t["need"] == "plan" else "plan"
            intentos = [(t["need"], True), (otra, True), (t["need"], False), (otra, False)]
        desde_t = DESDE if t["tipo"] != "seguimiento" else sumar_habiles(M["metaJunaeb"], 1, M)
        for need, dentro in intentos:
            for d in dias:
                if d < desde_t or (dentro and d > t["deadline"] and t["tipo"] != "seguimiento"):
                    continue
                x = elegir_tec(t, d, need)
                if x:
                    lugar = (d, x, libres(d, x, need)[0])
                    break
            if lugar:
                break
        if not lugar:                     # sin cupo: bloque extra del primer día
            lugar = (dias[0], "RODRIGO", 4)
            avisos.append(f"Sin cupo para {t['rbd']} {E[t['rbd']]['nombre']}: quedó en bloque extra.")
        d, x, b = lugar
        ocup[(d, x)].add(b)
        sitios[(d, x)].append(coords(E[t["rbd"]], E, base_comuna))
        aplicar_trabajo(t, d, x, b, P, E, corr_de, nuevo_id, M, DESDE, reprog)

    D["PLAN"] = P
    for c in CORR:
        if c["quien"] == "tecnico":
            prox = sorted((p for p in P if p["rbd"] == c["rbd"] and p["estado"] == "programada"),
                          key=lambda p: p["fecha"])
            c["agendado"] = prox[0]["fecha"] if prox else None
        elif c["rbd"] == 995441:
            c["agendado"] = "2026-11-16"
    CORR.sort(key=lambda c: -c["pts"])
    D["CORR"] = CORR

    # bloques ordenados en los días ya pasados (las realizadas que se sumaron)
    for (d, tec) in {(p["fecha"], p["tec"]) for p in P if p["origen"] == "plan" and p["fecha"] < DESDE}:
        ts = [p for p in P if p["fecha"] == d and p["tec"] == tec and p["estado"] != "anulada"]
        ts.sort(key=lambda p: (not (p["clase"].startswith("VISITA") or p["clase"].startswith("PREVENTIVA")),
                               p["bloque"], p["id"]))
        mx = 2 if vie(d) else 3
        for i, p in enumerate(ts, 1):
            p["bloque"] = i if i <= mx else 4

    recalcular_fichas(D, E, CORR, M)
    for c in CORR:
        c["pts"] = E[c["rbd"]]["pts"]
    D["CORR"].sort(key=lambda c: -c["pts"])

    # ---------------------------------------------------------------- META
    M["hoy"] = hoy
    M["datacoraCorte"] = CORTE
    n = int((re.search(r"-r(\d+)$", version_entrada) or [0, 0])[1])
    M["version"] = a.version or f"{hoy}-r{n + 1}"
    M["firmasPrevias"] = {k: v for k, v in previas.items() if k != M["version"]}
    M["revisadoWsp"] = CORTE
    guardar_js(a.datos, D, F(hoy).strftime("%d-%m-%Y"))

    # ---------------------------------------------------------------- reporte
    okJ = {d["rbd"] for d in D["DC"] if d["cuenta"]} | \
          {p["rbd"] for p in P if p["estado"] == "realizada" and p["fecha"] >= CORTE}
    jun = [e for e in D["ESTAB"] if e["inst"] == "Junaeb"]
    jar = [e for e in D["ESTAB"] if e["inst"] != "Junaeb"]
    rep.append(f"Meta JUNAEB: {sum(1 for e in jun if e['rbd'] in okJ)} de {len(jun)} con visita en Datácora.")
    rep.append(f"Meta jardines: {sum(1 for e in jar if e['prevOK'])} de {len(jar)} con preventiva completa.")
    rep.append(f"Reprogramadas: {len(reprog)} tarjetas que no se ejecutaron en su fecha.")
    if avisos:
        rep.append("Avisos:")
        rep += ["  " + x for x in avisos]
    texto = "\n".join(rep)
    print(texto)
    if a.reporte:
        open(a.reporte, "w", encoding="utf-8").write(texto + "\n")


def firmas(D):
    """id -> 'estado|fecha|tec|bloque|aplaz|folio', lo mismo que compara fusionar() en el panel."""
    return {p["id"]: "|".join(str(x) for x in (p.get("estado", ""), p.get("fecha", ""), p.get("tec", ""),
                                                 p.get("bloque", ""), p.get("aplaz") or 0, p.get("folio") or ""))
            for p in D["PLAN"]}


def sumar_habiles(desde, n, M):
    hab = M["habiles"]
    d = F(desde)
    k = 0
    for _ in range(80):
        d += dt.timedelta(days=1)
        s = d.isoformat()
        if d.weekday() < 5 and s not in M["feriados"]:
            k += 1
            if k == n:
                return s
    return hab[-1]


def aplicar_trabajo(t, d, tec, b, P, E, corr_de, nuevo_id, M, DESDE, reprog):
    """Deja la tarjeta del trabajo en su nueva fecha, con un texto que se entienda en terreno."""
    e = E[t["rbd"]]
    c = corr_de.get(t["rbd"])
    p = t["card"]
    partes = []
    if p is None:
        p = {"id": nuevo_id("P"), "fase": "F1", "fecha": d, "tec": tec, "bloque": b, "hora": "",
             "rbd": t["rbd"], "clase": "", "detalle": "", "tipoReal": "", "pts": e["pts"], "crit": "",
             "estado": "programada", "aplaz": 0, "origen": "plan"}
        P.append(p)
    elif p["fecha"] < DESDE:
        p["aplaz"] = (p.get("aplaz") or 0) + 1
        reprog.append(p["id"])
        partes.append(f"Reprogramada: no se ejecutó el {fc(p['fecha'])}")
    p["fecha"], p["tec"], p["bloque"] = d, tec, b
    p["crit"] = c["crit"] if (c and c["quien"] == "tecnico") else ""
    p["pts"] = min(20, e["pts"] + (p.get("aplaz") or 0))
    det_corr = ("CORRECTIVO: " + c["detalle"]) if (c and c["quien"] == "tecnico") else ""
    meta = "VISITA para Datácora — este JUNAEB todavía no tiene bitácora en Datácora (meta al 30-10)"
    if t["tipo"] == "visita":
        p["clase"], p["tipoReal"] = "VISITA", "Visita"
        partes.append(meta)
    elif t["tipo"] == "visita+corr":
        p["clase"], p["tipoReal"] = "VISITA + CORRECTIVO", "Visita"
        partes += [meta, det_corr]
    elif t["tipo"] == "prev+corr":
        p["clase"], p["tipoReal"] = "PREVENTIVA + CORRECTIVO", "Plan Preventivo"
        falta = p.get("parte") or "JI"
        txt = {"SC": "falta la bitácora SC", "JI+SC": "2 bitácoras: JI y SC"}.get(falta, "bitácora JI")
        partes += [f"PREVENTIVA obligatoria ({txt}), adelantada para cerrar en la misma visita lo que pidió supervisión",
                   det_corr]
    elif t["tipo"] == "corr":
        p["clase"], p["tipoReal"] = "CORRECTIVO URGENTE", "Emergencia"
        partes.append(det_corr)
    elif t["tipo"] == "seguimiento":
        p["clase"], p["tipoReal"] = "VISITA", "Visita"
        partes.append("Visita de seguimiento: la meta JUNAEB ya está cumplida en Datácora; va después del 30-10")
    p["detalle"] = "  ||  ".join(x for x in partes if x)
    for o in t["absorbe"]:
        o["estado"] = "anulada"
        o["absorbidaPor"] = p["id"]
        o["detalle"] = f"Se resuelve en la visita del {fc(d)} ({p['id']})  ||  " + o["detalle"]


def recalcular_fichas(D, E, CORR, M):
    """Cobertura, última visita y puntaje de cada establecimiento con los pesos del README."""
    W = M["pesos"]
    S2 = M["corteS2"]
    corr = {c["rbd"]: c for c in CORR if c["quien"] == "tecnico"}
    real = [p for p in D["PLAN"] if p["estado"] == "realizada"]
    dc_por = defaultdict(list)
    for d in D["DC"]:
        dc_por[d["rbd"]].append(d)
    for e in D["ESTAB"]:
        r = e["rbd"]
        v2 = {(p["fecha"]) for p in real if p["rbd"] == r and p["fecha"] >= S2}
        v2 |= {h["fecha"] for h in D["HIST"] if h["rbd"] == r and h["fecha"] >= S2}
        v2 |= {d["fecha"] for d in dc_por[r]}
        corr2 = {p["fecha"] for p in real if p["rbd"] == r and p["fecha"] >= S2 and not es_prev(p.get("tipoReal"))}
        corr2 |= {h["fecha"] for h in D["HIST"] if h["rbd"] == r and h["fecha"] >= S2 and h["tipo"] != "Preventiva"}
        partes, fprev = set(), []
        for p in real:
            if p["rbd"] == r and p["fecha"] >= S2 and es_prev(p.get("tipoReal")):
                partes |= set((p.get("parte") or "JI").split("+"))
                fprev.append(p["fecha"])
        for d in dc_por[r]:
            if d["tipo"] == "Preventiva":
                partes.add("JI")
                fprev.append(d["fecha"])
        e["s2"] = len(v2)
        e["corr2"] = len(corr2)
        todas = sorted(v2 | {p["fecha"] for p in real if p["rbd"] == r})
        e["ultima"] = todas[-1] if todas else None
        e["prevPartes"] = sorted(partes)
        e["prevFecha"] = min(fprev) if fprev else None
        if e["inst"] == "Junaeb":
            e["prevOK"] = None
            completa = bool(partes)
        else:
            completa = ({"JI", "SC"} <= partes) if e["nbit"] == 2 else bool(partes)
            e["prevOK"] = completa
        e["enDatacora"] = any(d["cuenta"] for d in dc_por[r]) if e["inst"] == "Junaeb" else bool(dc_por[r])
        e["dcVisitas"] = len(dc_por[r])
        e["cobertura"] = "CON PREVENTIVA" if completa else ("SIN PREVENTIVA" if v2 else "SIN VISITA")
        # puntaje
        why, pts = [], 0

        def suma(txt, n):
            nonlocal pts
            if n:
                why.append(f"{txt} +{n}")
                pts += n
        if not v2:
            suma("sin visita 2°sem", W["sinVisita"])
            if not e.get("s1"):
                suma("sin visita en todo 2026", W["sinVisita2026"])
        elif not completa:
            suma("sin preventiva certificada" if e["inst"] == "Junaeb" else "preventiva incompleta", W["sinPreventiva"])
        c = corr.get(r)
        if c:
            n = W.get(c["crit"], 1)
            if c["estado"] == "PARCIAL":
                n = max(1, n // 2)
            suma(f"falla {c['crit']} abierta" + (" (parcial)" if c["estado"] == "PARCIAL" else ""), n)
            ant = min(3, c["dias"] // 15)
            suma(f"{c['dias']} días sin atender", ant * W["antiguedad"])
            if c["prio"]:
                suma("prioridad de supervisión", W["prioridad"])
        if (e.get("rac") or 0) >= 800:
            suma(f"{e['rac']} raciones", W["raciones800"])
        elif (e.get("rac") or 0) >= 400:
            suma(f"{e['rac']} raciones", W["raciones400"])
        if e["inst"] != "Junaeb":
            suma("jardín/sala cuna", W["jardin"])
        if re.search(r"granel|45k", e.get("gas") or "", re.I):
            suma(f"gas {e['gas']}", W["gasRiesgo"])
        e["pts"] = max(1, min(20, pts))
        e["why"] = "; ".join(why) or "sin factores de riesgo"


if __name__ == "__main__":
    main()
